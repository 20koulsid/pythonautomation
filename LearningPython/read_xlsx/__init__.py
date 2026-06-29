from openpyxl import Workbook, load_workbook
wb = load_workbook(filename='student_record.xlsx')
sh = wb['studentRecord']
print(sh['A3'].value)
print(sh['A5'].value)
# OR
print(wb['studentRecord']['A3'].value)
print(sh.cell(row=1,column=3).value)

row_ct = sh.max_row
col_ct = sh.max_column
print(row_ct, col_ct)


# for n number of data we can write below code for row-i, column-j
for i in range(1,row_ct+1):
    for j in range(1,col_ct+1):
        print(sh.cell(row=i,column=j).value,end=" ")
        print('\n')
