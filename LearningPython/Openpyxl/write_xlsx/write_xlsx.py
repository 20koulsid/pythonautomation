# # Openpyxl is a Python library to read/write Excel 2010 xlsx/xlsm/xltx/xltm files.
import os
import openpyxl
from openpyxl import Workbook
#create a workbook
wb = Workbook()

# grab the active worksheet
ws = wb.active
#
# # Data can be assigned directly to cells
# ws['A1'] = "My name is Sid"
#
# # Rows can also be appended
# ws.append([1, 2, 3])
#
# wb.save('demoxl.xlsx')

# testdata = [['Name','City'],['vinchester','New York'],['russia','Russia'],['desposa','triadent']]
# for data in testdata:
#     ws.append(data)

# wb.save('test.xlsx')
# os.startfile('test.xlsx') # this os will open excel directly
for i in range(1,6): #this is for row
    for j in range(1,5): #this is for column
        ws.cell(row=i,column=j).value = i+j
wb.save('demo.xlsx')
os.startfile('demo.xlsx')
