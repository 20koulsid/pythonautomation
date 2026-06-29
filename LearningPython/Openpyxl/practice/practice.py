# 🟡 Beginner Level Questions
# Create a new Excel workbook and add a sheet named Students.
import os
from os import name

import os
from openpyxl import Workbook

wb = Workbook()
ws = wb.active

# ws.title = "Students"
#
# # header
# ws.append(["Name"])
#
# # data rows
# testdata = ["Rimal", "Sid", "Champak", "Sunita"]
#
# for name in testdata:
#     ws.append([name])
#
# wb.save("student.xlsx")
#
# os.startfile("student.xlsx")
# # Add column headers: Name, Age, City in the first row.
# ws.append(['Name','Age','City'])
# wb.save("student.xlsx")
# os.startfile("student.xlsx")
# Insert at least 5 student records into the sheet.
student_data = [['Name','Class','marks'],['Rahul','10th','230'],['Yash','10th','300'],['Rahil','10th','300']
                ,['Ronit','10th','220'],['Rajesh','10th','100']]

for data in student_data:
    ws.append(data)
wb.save("student_record.xlsx")
os.startfile("student_record.xlsx")
# Save the workbook as students.xlsx.
# Open an existing Excel file and print all rows one by one.
from openpyxl import load_workbook

# Load existing file
wb = load_workbook("student_record.xlsx")
ws = wb.active

# Print all rows one by one
for row in ws.iter_rows(values_only=True):
    print(row)
# Read only the “Name” column from an Excel sheet.
# Append 3 new rows to an existing Excel file.
# Create a workbook with 3 sheets named Math, Science, and English.
# 🟡 Intermediate Level Questions
# Update the city of a specific student in the Excel file.
# Find and print the student with the highest marks.
# Filter and display students who scored more than 75 marks.
# Delete rows where marks are below 40.
# Count the total number of students in the sheet.
# Find duplicate names in the Excel sheet.
# Update multiple cells using row and column indexing.
# Copy data from one sheet to another sheet.
# 🔵 Advanced Level Questions
# Generate a student report with Total Marks, Average Marks, and Grade.
# Apply grading system:
# A → >85
# B → 60–85
# C → <60
# Format Excel headers (bold, font size, background color).
# Auto-adjust column widths for all columns.
# Merge cells and create a title “STUDENT REPORT”.
# Create an attendance sheet and calculate attendance percentage.
# Build a salary sheet with Basic, Bonus, and Total Salary.
# Create a summary sheet combining data from multiple sheets.
# Generate an Excel dashboard with total counts and averages.
# 🏆 Challenge Level Questions
# Build a mini Student Management System using Excel:
# Add student
# Update student
# Delete student
# Search student
# Create an automated expense tracker in Excel.
# Generate monthly reports from raw data automatically.
# Build a multi-sheet employee management system.
# Create a dynamic report generator that updates data automatically when new rows are added.