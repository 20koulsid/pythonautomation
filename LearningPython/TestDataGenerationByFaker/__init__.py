import os

from faker import Faker
from faker.providers.date_time import hi_IN
from openpyxl import Workbook
wb = Workbook()
ws = wb.active

fake_data = Faker(['hi_IN'])
# .name/.email/.fullname/.address and so on will everytime create a new data
# print(fake_data.name())
for i in range(1,13):
    for j in range(1,3):
        ws.cell(row=i,column=1).value = fake_data.name()
        ws.cell(row=i, column=2).value = fake_data.email()
        ws.cell(row=i, column=3).value = fake_data.address()
        wb.save("test.xlsx")
os.startfile("test.xlsx")

# inorder to add names in any other languge or anything we need to add local
# Faker(['it_IT'])
